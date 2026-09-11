# -*- coding: utf-8 -*-
"""test_extractors — 类型检测 / YouTube / B站 / 网页 / 本地提取器"""
import json
import types

import pytest

import extractors as ex


# ---------- 类型检测 ----------

def test_detect_urls():
    assert ex.detect_content_type("https://www.youtube.com/watch?v=abc") == "youtube"
    assert ex.detect_content_type("https://youtu.be/abcdefghijk") == "youtube"
    assert ex.detect_content_type("https://www.bilibili.com/video/BV1xx") == "bilibili"
    assert ex.detect_content_type("https://b23.tv/BV1xx") == "bilibili"
    assert ex.detect_content_type("https://example.com/a") == "web"
    assert ex.detect_content_type(None) == "unknown"


def test_detect_files(tmp_path):
    cases = {".txt": "text", ".md": "text", ".csv": "text", ".pdf": "pdf",
             ".docx": "word", ".doc": "word", ".epub": "epub", ".xlsx": "excel",
             ".xlsm": "excel", ".pptx": "pptx", ".mp3": "audio", ".wav": "audio",
             ".mp4": "video", ".mkv": "video"}
    for ext, kind in cases.items():
        p = tmp_path / f"f{ext}"
        p.write_text("x", encoding="utf-8")
        assert ex.detect_content_type(str(p)) == kind, ext
    assert ex.detect_content_type(str(tmp_path / "nope.xyz")) == "unknown"


def test_video_id_and_bvid():
    assert ex.extract_video_id("https://www.youtube.com/watch?v=dQw4w9WgXcQ") == "dQw4w9WgXcQ"
    assert ex.extract_video_id("https://youtu.be/dQw4w9WgXcQ") == "dQw4w9WgXcQ"
    assert ex.extract_video_id("https://example.com") is None
    assert ex.extract_bvid("https://www.bilibili.com/video/BV1GJ411x7h7") == "BV1GJ411x7h7"


def test_clean_subtitle():
    raw = "1\n00:00:01,000 --> 00:00:02,000\n<b>你好</b>\n\nWEBVTT\nKind: captions\n"
    out = ex.clean_subtitle(raw)
    assert "00:00" not in out and "<b>" not in out and "WEBVTT" not in out and "你好" in out


# ---------- cookies 规则 ----------

def test_yt_cookies_missing_file(tmp_path, monkeypatch):
    monkeypatch.setenv("MYAGENTRAG_YOUTUBE_COOKIES", str(tmp_path / "nope.txt"))
    assert ex._yt_cookies_args() == []


def test_yt_cookies_empty_file_ignored(tmp_path, monkeypatch):
    p = tmp_path / "c.txt"
    p.write_text("", encoding="utf-8")
    monkeypatch.setenv("MYAGENTRAG_YOUTUBE_COOKIES", str(p))
    assert ex._yt_cookies_args() == []


def test_bilibili_cookie_header_netscape(tmp_path, monkeypatch):
    p = tmp_path / "bili.txt"
    p.write_text("# Netscape HTTP Cookie File\n"
                 ".bilibili.com\tTRUE\t/\tTRUE\t0\tSESSDATA\tsess_value\n"
                 ".bilibili.com\tTRUE\t/\tTRUE\t0\tbili_jct\tjct_value\n"
                 ".example.com\tTRUE\t/\tTRUE\t0\tOTHER\tx\n", encoding="utf-8")
    monkeypatch.delenv("MYAGENTRAG_BILIBILI_COOKIES", raising=False)
    monkeypatch.setattr(ex, "MANAGED_HOME", tmp_path / "managed")
    # _bilibili_cookie_header 默认读 MANAGED_HOME/cookies/...，直接传环境变量更稳
    monkeypatch.setenv("MYAGENTRAG_BILIBILI_COOKIES", str(p))
    hdr = ex._bilibili_cookie_header()
    assert "SESSDATA=sess_value" in hdr and "bili_jct=jct_value" in hdr and "OTHER" not in hdr


def test_bilibili_cookie_header_raw(tmp_path, monkeypatch):
    p = tmp_path / "bili2.txt"
    p.write_text("SESSDATA=abc; bili_jct=def\n", encoding="utf-8")
    monkeypatch.setenv("MYAGENTRAG_BILIBILI_COOKIES", str(p))
    assert ex._bilibili_cookie_header() == "SESSDATA=abc; bili_jct=def"


# ---------- 文本 / Excel ----------

def test_extract_text_file_encodings(tmp_path):
    u = tmp_path / "u.txt"
    u.write_text("中文内容", encoding="utf-8")
    assert ex.extract_text_file(str(u)) == "中文内容"
    g = tmp_path / "g.txt"
    g.write_bytes("中文内容".encode("gbk"))
    assert ex.extract_text_file(str(g)) == "中文内容"


def test_extract_excel_text(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    from openpyxl import Workbook
    p = tmp_path / "t.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "表一"
    ws.append(["k", "v"])
    ws.append(["a", "b"])
    wb.save(p)
    out = ex.extract_excel_text(str(p))
    assert "## 工作表: 表一" in out and "a | b" in out


# ---------- YouTube：元数据解析回归（防 import json 缺失复发）+ 原始字幕捕获 ----------

def test_extract_youtube_metadata_and_raw_subtitle(tmp_path, monkeypatch):
    vtt = """WEBVTT

00:00:00.000 --> 00:00:02.000
Hello world

00:00:02.000 --> 00:00:04.000
Second line
"""
    info = json.dumps({"title": "测试视频", "uploader": "测试UP"})
    calls = {"n": 0}

    def fake_run(cmd, **kw):
        calls["n"] += 1
        return types.SimpleNamespace(returncode=0, stdout=info if calls["n"] == 1 else "", stderr="")

    monkeypatch.setattr(ex, "_find_ytdlp", lambda: "yt-dlp")
    monkeypatch.setattr(ex, "_yt_cookies_args", lambda: [])
    monkeypatch.setattr(ex, "make_tmpdir", lambda prefix: tmp_path)
    (tmp_path / "dQw4w9WgXcQ.en.vtt").write_text(vtt, encoding="utf-8")
    monkeypatch.setattr(ex.subprocess, "run", fake_run)

    r = ex.extract_youtube("dQw4w9WgXcQ")
    assert r["success"] is True
    assert r["title"] == "测试视频" and r["author"] == "测试UP"
    assert r["raw_subtitle"].startswith("WEBVTT") and r["raw_subtitle_ext"] == "vtt"
    assert "Hello world" in r["transcript"]


def test_extract_youtube_no_ytdlp(monkeypatch):
    monkeypatch.setattr(ex, "_find_ytdlp", lambda: None)
    r = ex.extract_youtube("dQw4w9WgXcQ")
    assert r["success"] is False and "yt-dlp" in r["error"]


# ---------- B站：API 链路（requests 打桩） ----------

class _FakeResp:
    def __init__(self, payload):
        self._p = payload

    def json(self):
        return self._p


def test_extract_bilibili_with_subtitles(monkeypatch):
    view = {"code": 0, "data": {"title": "B站标题", "owner": {"name": "UP主"}, "cid": "777"}}
    player = {"code": 0, "data": {"subtitle": {"subtitles": [
        {"subtitle_url": "//api.bilibili.com/x/sub.json"}]}}}
    sub = {"body": [{"from": 0.5, "to": 2.5, "content": "第一句"},
                    {"from": 3.0, "to": 5.0, "content": "第二句"}]}

    def fake_get(url, **kw):
        if "web-interface/view" in url:
            return _FakeResp(view)
        if "player/wbi" in url:
            return _FakeResp(player)
        return _FakeResp(sub)

    fake_requests = types.SimpleNamespace(get=fake_get)
    monkeypatch.setitem(__import__("sys").modules, "requests", fake_requests)
    monkeypatch.setattr(ex, "_bilibili_cookie_header", lambda: "")
    r = ex.extract_bilibili("BV1fake00000")
    assert r["success"] is True
    assert r["title"] == "B站标题" and r["author"] == "UP主"
    assert len(r["subtitle_segments"]) == 2
    assert r["subtitle_segments"][0] == {"start_ms": 500, "end_ms": 2500, "text": "第一句"}
    assert r["raw_subtitle_ext"] == "json"
    assert "第一句" in r["transcript"]


def test_extract_bilibili_no_subtitles(monkeypatch):
    view = {"code": 0, "data": {"title": "无字幕", "owner": {"name": "UP"}, "cid": "1"}}
    player = {"code": 0, "data": {"subtitle": {"subtitles": []}}}

    def fake_get(url, **kw):
        return _FakeResp(view) if "web-interface/view" in url else _FakeResp(player)

    monkeypatch.setitem(__import__("sys").modules, "requests", types.SimpleNamespace(get=fake_get))
    monkeypatch.setattr(ex, "_bilibili_cookie_header", lambda: "")
    r = ex.extract_bilibili("BV1fake00001")
    assert r["success"] is False and r["title"] == "无字幕"  # 无 CC 字幕 success:false 属正常


# ---------- EPUB：常量模块路径回归（端侧验证发现，ebooklib 0.20） ----------

def test_extract_epub_text(tmp_path):
    """回归：ITEM_DOCUMENT 等常量必须取自顶层 ebooklib 模块。
    ebooklib 0.20 起不再把常量暴露到 ebooklib.epub 命名空间，
    旧写法 epub.ITEM_DOCUMENT 会 AttributeError → EPUB 提取整体失败。"""
    pytest.importorskip("ebooklib")
    from ebooklib import epub as epub_mod

    book = epub_mod.EpubBook()
    book.set_identifier("myag-test-id")
    book.set_title("测试电子书")
    book.set_language("zh")
    chapter = epub_mod.EpubHtml(title="第一章", file_name="chap_01.xhtml", lang="zh")
    chapter.content = "<html><body><p>第一章正文内容标记。</p></body></html>"
    book.add_item(chapter)
    book.add_item(epub_mod.EpubNcx())
    book.add_item(epub_mod.EpubNav())
    book.spine = ["nav", chapter]
    p = tmp_path / "t.epub"
    epub_mod.write_epub(str(p), book)

    out = ex.extract_epub_text(str(p))
    assert isinstance(out, ex.Extraction)  # v0.8.0 §8D：Extraction{text, srcmap}
    assert "第一章正文内容标记" in out.text


def test_extract_epub_headings_and_chapters(tmp_path):
    """§8D：h 标签归一化为 markdown # 标记；spine 章节账本带偏移与标题"""
    pytest.importorskip("ebooklib")
    from ebooklib import epub as epub_mod

    book = epub_mod.EpubBook()
    book.set_identifier("myag-test-h2")
    book.set_title("测试电子书2")
    book.set_language("zh")
    ch1 = epub_mod.EpubHtml(title="第一章", file_name="chap_01.xhtml", lang="zh")
    ch1.content = ('<html><body><h1>第一章 标题甲</h1><p>第一章正文内容。</p>'
                   '<h2>第一节 小标题</h2><p>小节内容文字。</p></body></html>')
    ch2 = epub_mod.EpubHtml(title="第二章", file_name="chap_02.xhtml", lang="zh")
    ch2.content = '<html><body><p>第二章正文内容。</p></body></html>'
    book.add_item(ch1)
    book.add_item(ch2)
    book.add_item(epub_mod.EpubNcx())
    book.add_item(epub_mod.EpubNav())
    book.spine = ["nav", ch1, ch2]
    book.toc = (epub_mod.Section("目录"),
                epub_mod.Link("chap_01.xhtml", "第一章 标题甲", "c1"),
                epub_mod.Link("chap_02.xhtml", "第二章 标题乙", "c2"))
    p = tmp_path / "t2.epub"
    epub_mod.write_epub(str(p), book)

    out = ex.extract_epub_text(str(p))
    assert isinstance(out, ex.Extraction)
    assert "# 第一章 标题甲" in out.text
    assert "## 第一节 小标题" in out.text
    assert out.srcmap and out.srcmap["kind"] == "epub"
    chs = out.srcmap["chapters"]
    assert len(chs) == 2
    assert chs[0][0] == 0 and chs[1][0] > 0        # 章节起始偏移递增
    assert chs[0][2] == "第一章 标题甲"              # toc 标题映射


# ---------- 网页 ----------

def test_extract_web(monkeypatch):
    jina = "Title: 测试页面\n\n正文第一行\n正文第二行"

    class R:
        status_code = 200
        text = jina

    monkeypatch.setitem(__import__("sys").modules, "requests",
                        types.SimpleNamespace(get=lambda url, **kw: R()))
    r = ex.extract_web("https://example.com/a")
    assert r["success"] is True
    assert r["title"] == "测试页面" and "正文第一行" in r["content"]


def test_extract_web_failure(monkeypatch):
    class R:
        status_code = 500
        text = ""
    monkeypatch.setitem(__import__("sys").modules, "requests",
                        types.SimpleNamespace(get=lambda url, **kw: R()))
    r = ex.extract_web("https://example.com/a")
    assert r["success"] is False


# ---------- §8D 结构化提取：Extraction/srcmap ----------

def test_extract_pdf_page_ledger(tmp_path):
    """§8D：PDF 逐页偏移记账——outline 锚点落库的硬前提"""
    fitz = pytest.importorskip("pymupdf")
    doc = fitz.open()
    texts = ["page one alpha beta", "page two gamma delta"]  # 默认字体不支持 CJK
    for t in texts:
        page = doc.new_page()
        page.insert_text((72, 72), t)
    p = tmp_path / "t.pdf"
    doc.save(str(p))
    doc.close()
    out = ex.extract_pdf_text(str(p))
    assert isinstance(out, ex.Extraction)
    pages = out.srcmap["pages"]
    assert pages[0] == [0, 1] and pages[1][1] == 2
    assert pages[1][0] == len(out.text.split("\n")[0]) + 1  # 账本与 full.md 坐标一致
    assert "page two gamma delta" in out.text


def test_extract_docx_heading_markers(tmp_path):
    """§8D：docx Heading 样式 → markdown # 标记归一化"""
    pytest.importorskip("docx")
    from docx import Document
    d = Document()
    d.add_heading("第一章 总览", level=1)
    d.add_paragraph("正文段落内容标记。")
    d.add_heading("第二节 细节", level=2)
    d.add_paragraph("第二段正文内容标记。")
    p = tmp_path / "t.docx"
    d.save(str(p))
    out = ex.extract_word_text(str(p))
    assert isinstance(out, ex.Extraction)
    assert "# 第一章 总览" in out.text and "## 第二节 细节" in out.text
    assert "正文段落内容标记。" in out.text


def test_pdf_noise_filter():
    """页脚噪声过滤：签名行删除、正文行保留"""
    page = ("正文内容保持。甲乙丙\n"
            "file:///D|/C:/book/item_14.html (1 of 3)2005-4-26 15:16:19\n"
            "(2 of 3)2005-4-26 15:16:19\n"
            "2005-4-26 15:16:19\n"
            "正文继续内容丁戊。")
    out = ex._pdf_noise_filter(page)
    assert "file:///" not in out and "正文继续内容丁戊。" in out and "甲乙丙" in out


# ---------- OPT-06：--url 抓取失败结构化（QA 2026-09-11） ----------

def test_extract_web_non200_structured(monkeypatch):
    """OPT-06：非 200 不再静默——error + error_i18n 双语（含状态码与 URL）"""
    import sys, types
    import extractors

    class FakeResp:
        status_code = 404
        text = ""

    fake = types.SimpleNamespace(get=lambda url, timeout=60: FakeResp())
    monkeypatch.setitem(sys.modules, "requests", fake)
    r = extractors.extract_web("https://example.com/missing")
    assert r["success"] is False
    assert "404" in r["error"] and "example.com/missing" in r["error"]
    assert r["error_i18n"]["en"].startswith("Web fetch failed (HTTP 404)")


def test_extract_web_exception_bilingual(monkeypatch):
    """OPT-06 附带：网络异常路径补 error_i18n 双语（原仅裸 str(e)）"""
    import sys, types
    import extractors

    def boom(url, timeout=60):
        raise ConnectionError("网络不可达")

    fake = types.SimpleNamespace(get=boom)
    monkeypatch.setitem(sys.modules, "requests", fake)
    r = extractors.extract_web("https://example.com/x")
    assert r["success"] is False
    assert "网络不可达" in r["error"]
    assert "error_i18n" in r and "Web fetch error" in r["error_i18n"]["en"]
